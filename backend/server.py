from dotenv import load_dotenv
load_dotenv()

from fastapi import FastAPI, APIRouter, HTTPException, Request, Response
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
from bson import ObjectId
import os
import logging
import bcrypt
import jwt as pyjwt
import json
import uuid
import asyncio
from pathlib import Path
from pydantic import BaseModel
from typing import List, Optional
from datetime import datetime, timezone, timedelta
from fastapi.responses import FileResponse
from openai import AsyncOpenAI
from pptx import Presentation as PptxPresentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from fpdf import FPDF

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

ROOT_DIR = Path(__file__).parent
GENERATED_DIR = ROOT_DIR / "generated"
GENERATED_DIR.mkdir(exist_ok=True)

# MongoDB
mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

# JWT
JWT_ALGORITHM = "HS256"

def get_jwt_secret():
    return os.environ["JWT_SECRET"]

# ── Password utils ──
def hash_password(password: str) -> str:
    return bcrypt.hashpw(password.encode("utf-8"), bcrypt.gensalt()).decode("utf-8")

def verify_password(plain: str, hashed: str) -> bool:
    return bcrypt.checkpw(plain.encode("utf-8"), hashed.encode("utf-8"))

# ── Token utils ──
def create_access_token(user_id: str, email: str) -> str:
    payload = {"sub": user_id, "email": email, "exp": datetime.now(timezone.utc) + timedelta(minutes=60), "type": "access"}
    return pyjwt.encode(payload, get_jwt_secret(), algorithm=JWT_ALGORITHM)

def create_refresh_token(user_id: str) -> str:
    payload = {"sub": user_id, "exp": datetime.now(timezone.utc) + timedelta(days=7), "type": "refresh"}
    return pyjwt.encode(payload, get_jwt_secret(), algorithm=JWT_ALGORITHM)

def set_auth_cookies(response: Response, access_token: str, refresh_token: str):
    response.set_cookie(key="access_token", value=access_token, httponly=True, secure=False, samesite="lax", max_age=3600, path="/")
    response.set_cookie(key="refresh_token", value=refresh_token, httponly=True, secure=False, samesite="lax", max_age=604800, path="/")

# ── Auth helper ──
async def get_current_user(request: Request) -> dict:
    token = request.cookies.get("access_token")
    if not token:
        auth = request.headers.get("Authorization", "")
        if auth.startswith("Bearer "):
            token = auth[7:]
    if not token:
        raise HTTPException(401, "Not authenticated")
    try:
        payload = pyjwt.decode(token, get_jwt_secret(), algorithms=[JWT_ALGORITHM])
        if payload.get("type") != "access":
            raise HTTPException(401, "Invalid token type")
        user = await db.users.find_one({"_id": ObjectId(payload["sub"])})
        if not user:
            raise HTTPException(401, "User not found")
        user["_id"] = str(user["_id"])
        user.pop("password_hash", None)
        return user
    except pyjwt.ExpiredSignatureError:
        raise HTTPException(401, "Token expired")
    except pyjwt.InvalidTokenError:
        raise HTTPException(401, "Invalid token")

# ── Brute force protection ──
async def check_brute_force(ip: str, email: str):
    identifier = f"{ip}:{email}"
    attempt = await db.login_attempts.find_one({"identifier": identifier}, {"_id": 0})
    if attempt and attempt.get("locked_until"):
        locked = datetime.fromisoformat(attempt["locked_until"]) if isinstance(attempt["locked_until"], str) else attempt["locked_until"]
        if locked > datetime.now(timezone.utc):
            raise HTTPException(429, "Too many failed attempts. Try again in 15 minutes.")
        else:
            await db.login_attempts.delete_one({"identifier": identifier})

async def record_failed_attempt(ip: str, email: str):
    identifier = f"{ip}:{email}"
    attempt = await db.login_attempts.find_one({"identifier": identifier}, {"_id": 0})
    if attempt:
        new_count = attempt.get("attempts", 0) + 1
        update = {"$set": {"attempts": new_count, "last_attempt": datetime.now(timezone.utc).isoformat()}}
        if new_count >= 5:
            update["$set"]["locked_until"] = (datetime.now(timezone.utc) + timedelta(minutes=15)).isoformat()
        await db.login_attempts.update_one({"identifier": identifier}, update)
    else:
        await db.login_attempts.insert_one({"identifier": identifier, "attempts": 1, "last_attempt": datetime.now(timezone.utc).isoformat()})

async def clear_failed_attempts(ip: str, email: str):
    await db.login_attempts.delete_one({"identifier": f"{ip}:{email}"})

# ── LLM Clients ──
def get_llm_client(role: str):
    configs = {
        "orchestrator": ("ORCHESTRATOR_API_KEY", "ORCHESTRATOR_BASE_URL", "ORCHESTRATOR_MODEL"),
        "content": ("CONTENT_API_KEY", "CONTENT_BASE_URL", "CONTENT_MODEL"),
        "layout": ("LAYOUT_API_KEY", "LAYOUT_BASE_URL", "LAYOUT_MODEL"),
    }
    key_env, url_env, model_env = configs[role]
    api_key = os.environ.get(key_env, "")
    base_url = os.environ.get(url_env, "")
    model = os.environ.get(model_env, "")
    if not api_key:
        return None, model
    return AsyncOpenAI(api_key=api_key, base_url=base_url), model

# ── Mock generators (fallback when no LLM keys) ──
def mock_orchestrate(prompt: str) -> dict:
    words = prompt.strip().split()
    title = prompt[:80] if len(prompt) <= 80 else " ".join(words[:8]) + "..."
    return {
        "title": title,
        "slides": [
            {"title": title, "type": "title"},
            {"title": "The Problem", "type": "content"},
            {"title": "Our Solution", "type": "content"},
            {"title": "Key Features", "type": "content"},
            {"title": "Business Model", "type": "content"},
            {"title": "Market Opportunity", "type": "content"},
            {"title": "Conclusion & Next Steps", "type": "content"},
        ]
    }

def mock_content_for_slide(slide_title: str, prompt: str) -> dict:
    templates = {
        "The Problem": [
            "Current solutions are inefficient and costly",
            "Users face significant friction in their workflow",
            "Existing tools lack integration and modern UX",
            "Market demand is growing rapidly with no clear leader",
        ],
        "Our Solution": [
            "AI-powered platform that automates key workflows",
            "Intuitive interface designed for speed and simplicity",
            "Seamless integration with existing tools and systems",
            "Built for scalability from day one",
        ],
        "Key Features": [
            "Intelligent automation that saves hours per week",
            "Real-time collaboration across teams",
            "Advanced analytics and reporting dashboard",
            "Enterprise-grade security and compliance",
        ],
        "Business Model": [
            "Freemium model with tiered enterprise pricing",
            "Monthly recurring revenue with annual discount",
            "Target: $10-50/user/month depending on tier",
            "Expansion revenue from add-on modules",
        ],
        "Market Opportunity": [
            "Total addressable market exceeds $50B globally",
            "Growing at 25% CAGR over the next 5 years",
            "Key segments: SMBs, mid-market, and enterprise",
            "First-mover advantage in our specific niche",
        ],
        "Conclusion & Next Steps": [
            "Strong product-market fit validated by early users",
            "Seeking strategic partnerships and funding",
            "Roadmap: expand features and enter new markets",
            "Join us in reshaping the future of this industry",
        ],
    }
    bullets = templates.get(slide_title, [
        f"Key insight about {slide_title.lower()}",
        f"Important detail regarding {slide_title.lower()}",
        f"Strategic consideration for {slide_title.lower()}",
        f"Action item related to {slide_title.lower()}",
    ])
    return {"title": slide_title, "bullets": bullets}

def mock_layout_for_slide(slide: dict) -> dict:
    slide_type = slide.get("type", "content")
    if slide_type == "title":
        return {**slide, "layout": "title_centered", "positions": {"title": {"x": 1.0, "y": 2.5, "w": 8.0, "h": 1.5}, "subtitle": {"x": 1.5, "y": 4.2, "w": 7.0, "h": 1.0}}}
    return {**slide, "layout": "title_bullets", "positions": {"title": {"x": 0.8, "y": 0.4, "w": 8.5, "h": 1.0}, "bullets": {"x": 1.0, "y": 1.8, "w": 8.0, "h": 4.5}}}

# ── LLM-powered generators ──
async def llm_orchestrate(prompt: str) -> dict:
    llm_client, model = get_llm_client("orchestrator")
    if not llm_client:
        return mock_orchestrate(prompt)
    try:
        resp = await llm_client.chat.completions.create(
            model=model,
            messages=[
                {"role": "system", "content": "You are a presentation outline generator. Given a topic/prompt, generate a structured presentation outline. Return ONLY valid JSON with this format: {\"title\": \"Presentation Title\", \"slides\": [{\"title\": \"Slide Title\", \"type\": \"title\"}, {\"title\": \"Slide Title\", \"type\": \"content\"}, ...]}. Generate 6-8 slides. First slide should be type 'title', rest should be 'content'."},
                {"role": "user", "content": f"Create a presentation outline for: {prompt}"}
            ],
            temperature=0.7,
            max_tokens=1000,
        )
        content = resp.choices[0].message.content.strip()
        if content.startswith("```"):
            content = content.split("```")[1]
            if content.startswith("json"):
                content = content[4:]
        return json.loads(content)
    except Exception as e:
        logger.error(f"Orchestrator LLM error: {e}")
        return mock_orchestrate(prompt)

async def llm_content_for_slide(slide_title: str, prompt: str) -> dict:
    llm_client, model = get_llm_client("content")
    if not llm_client:
        return mock_content_for_slide(slide_title, prompt)
    try:
        resp = await llm_client.chat.completions.create(
            model=model,
            messages=[
                {"role": "system", "content": "You generate slide content for presentations. Given a slide title and context, return ONLY valid JSON: {\"title\": \"Slide Title\", \"bullets\": [\"bullet 1\", \"bullet 2\", ...]}. Generate 3-5 concise, impactful bullet points. Keep each bullet under 80 characters."},
                {"role": "user", "content": f"Slide title: '{slide_title}'\nPresentation context: {prompt}\n\nGenerate bullet points for this slide."}
            ],
            temperature=0.7,
            max_tokens=500,
        )
        content = resp.choices[0].message.content.strip()
        if content.startswith("```"):
            content = content.split("```")[1]
            if content.startswith("json"):
                content = content[4:]
        return json.loads(content)
    except Exception as e:
        logger.error(f"Content LLM error: {e}")
        return mock_content_for_slide(slide_title, prompt)

async def llm_layout_for_slide(slide: dict) -> dict:
    llm_client, model = get_llm_client("layout")
    if not llm_client:
        return mock_layout_for_slide(slide)
    try:
        resp = await llm_client.chat.completions.create(
            model=model,
            messages=[
                {"role": "system", "content": "You assign layout positions for presentation slides. Given slide data, return ONLY valid JSON adding 'layout' and 'positions' fields. Layout types: 'title_centered' for title slides, 'title_bullets' for content slides. Positions use inches: {\"title\": {\"x\": float, \"y\": float, \"w\": float, \"h\": float}, \"bullets\": {\"x\": float, \"y\": float, \"w\": float, \"h\": float}}. Slide is 10x7.5 inches."},
                {"role": "user", "content": f"Assign layout for this slide: {json.dumps(slide)}"}
            ],
            temperature=0.3,
            max_tokens=300,
        )
        content = resp.choices[0].message.content.strip()
        if content.startswith("```"):
            content = content.split("```")[1]
            if content.startswith("json"):
                content = content[4:]
        return json.loads(content)
    except Exception as e:
        logger.error(f"Layout LLM error: {e}")
        return mock_layout_for_slide(slide)

# ── PPTX Renderer ──
def render_pptx(pres_id: str, slides: list, title: str) -> str:
    prs = PptxPresentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    for i, slide_data in enumerate(slides):
        sld = prs.slides.add_slide(blank_layout)
        layout_type = slide_data.get("layout", "title_bullets")
        positions = slide_data.get("positions", {})

        # Background
        background = sld.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        if layout_type == "title_centered":
            # Title slide
            tp = positions.get("title", {"x": 1.0, "y": 2.5, "w": 8.0, "h": 1.5})
            txBox = sld.shapes.add_textbox(Inches(tp["x"]), Inches(tp["y"]), Inches(tp["w"]), Inches(tp["h"]))
            tf = txBox.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = slide_data.get("title", title)
            p.font.size = Pt(36)
            p.font.bold = True
            p.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
            p.alignment = PP_ALIGN.CENTER

            # Subtitle
            sp = positions.get("subtitle", {"x": 1.5, "y": 4.2, "w": 7.0, "h": 1.0})
            txBox2 = sld.shapes.add_textbox(Inches(sp["x"]), Inches(sp["y"]), Inches(sp["w"]), Inches(sp["h"]))
            tf2 = txBox2.text_frame
            tf2.word_wrap = True
            p2 = tf2.paragraphs[0]
            p2.text = slide_data.get("bullets", [""])[0] if slide_data.get("bullets") else ""
            p2.font.size = Pt(18)
            p2.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
            p2.alignment = PP_ALIGN.CENTER

            # Blue accent line
            line_shape = sld.shapes.add_shape(1, Inches(3.5), Inches(4.0), Inches(3.0), Inches(0.04))
            line_shape.fill.solid()
            line_shape.fill.fore_color.rgb = RGBColor(0x00, 0x52, 0xFF)
            line_shape.line.fill.background()
        else:
            # Content slide - blue top bar
            bar = sld.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(0.06))
            bar.fill.solid()
            bar.fill.fore_color.rgb = RGBColor(0x00, 0x52, 0xFF)
            bar.line.fill.background()

            # Title
            tp = positions.get("title", {"x": 0.8, "y": 0.4, "w": 8.5, "h": 1.0})
            txBox = sld.shapes.add_textbox(Inches(tp["x"]), Inches(tp["y"]), Inches(tp["w"]), Inches(tp["h"]))
            tf = txBox.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = slide_data.get("title", "")
            p.font.size = Pt(28)
            p.font.bold = True
            p.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

            # Bullets
            bp = positions.get("bullets", {"x": 1.0, "y": 1.8, "w": 8.0, "h": 4.5})
            txBox2 = sld.shapes.add_textbox(Inches(bp["x"]), Inches(bp["y"]), Inches(bp["w"]), Inches(bp["h"]))
            tf2 = txBox2.text_frame
            tf2.word_wrap = True
            for j, bullet in enumerate(slide_data.get("bullets", [])):
                if j == 0:
                    para = tf2.paragraphs[0]
                else:
                    para = tf2.add_paragraph()
                para.text = bullet
                para.font.size = Pt(18)
                para.font.color.rgb = RGBColor(0x33, 0x40, 0x55)
                para.space_before = Pt(12)
                para.level = 0

            # Slide number
            num_box = sld.shapes.add_textbox(Inches(9.0), Inches(7.0), Inches(0.8), Inches(0.4))
            ntf = num_box.text_frame
            np = ntf.paragraphs[0]
            np.text = str(i + 1)
            np.font.size = Pt(10)
            np.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
            np.alignment = PP_ALIGN.RIGHT

    path = str(GENERATED_DIR / f"{pres_id}.pptx")
    prs.save(path)
    return path

# ── PDF Renderer ──
def render_pdf(pres_id: str, slides: list, title: str) -> str:
    pdf = FPDF(orientation='L', unit='mm', format='A4')
    pdf.set_auto_page_break(auto=False)

    for i, slide_data in enumerate(slides):
        pdf.add_page()
        layout_type = slide_data.get("layout", "title_bullets")

        if layout_type == "title_centered":
            pdf.set_fill_color(255, 255, 255)
            pdf.rect(0, 0, 297, 210, 'F')
            # Blue accent line
            pdf.set_fill_color(0, 82, 255)
            pdf.rect(108, 110, 80, 2, 'F')
            # Title
            pdf.set_font("Helvetica", "B", 32)
            pdf.set_text_color(15, 23, 42)
            pdf.set_xy(20, 60)
            pdf.cell(257, 20, slide_data.get("title", title), align='C')
            # Subtitle
            if slide_data.get("bullets"):
                pdf.set_font("Helvetica", "", 16)
                pdf.set_text_color(100, 116, 139)
                pdf.set_xy(30, 120)
                pdf.cell(237, 12, slide_data["bullets"][0][:80] if slide_data["bullets"] else "", align='C')
        else:
            pdf.set_fill_color(255, 255, 255)
            pdf.rect(0, 0, 297, 210, 'F')
            # Blue top bar
            pdf.set_fill_color(0, 82, 255)
            pdf.rect(0, 0, 297, 3, 'F')
            # Title
            pdf.set_font("Helvetica", "B", 24)
            pdf.set_text_color(15, 23, 42)
            pdf.set_xy(20, 15)
            pdf.cell(257, 14, slide_data.get("title", ""), align='L')
            # Bullets
            pdf.set_font("Helvetica", "", 16)
            pdf.set_text_color(51, 64, 85)
            y_pos = 45
            for bullet in slide_data.get("bullets", []):
                pdf.set_xy(30, y_pos)
                # Bullet marker
                pdf.set_fill_color(0, 82, 255)
                pdf.ellipse(26, y_pos + 3, 3, 3, 'F')
                pdf.cell(237, 10, bullet[:100], align='L')
                y_pos += 18
            # Slide number
            pdf.set_font("Helvetica", "", 9)
            pdf.set_text_color(148, 163, 184)
            pdf.set_xy(275, 195)
            pdf.cell(15, 8, str(i + 1), align='R')

    path = str(GENERATED_DIR / f"{pres_id}.pdf")
    pdf.output(path)
    return path

# ── Background generation task ──
async def generate_presentation_task(pres_id: str, prompt: str):
    try:
        # Step 1: Orchestrate
        logger.info(f"[{pres_id}] Step 1: Orchestrating outline...")
        outline = await llm_orchestrate(prompt)
        pres_title = outline.get("title", prompt[:80])
        slide_outlines = outline.get("slides", [])
        await db.presentations.update_one(
            {"id": pres_id},
            {"$set": {"status": "generating_content", "step": 2, "title": pres_title}}
        )

        # Step 2: Generate content for each slide
        logger.info(f"[{pres_id}] Step 2: Generating content for {len(slide_outlines)} slides...")
        slides_with_content = []
        for s in slide_outlines:
            s_title = s.get("title", "Untitled")
            s_type = s.get("type", "content")
            if s_type == "title":
                slides_with_content.append({"title": pres_title, "type": "title", "bullets": [prompt[:80]]})
            else:
                content = await llm_content_for_slide(s_title, prompt)
                content["type"] = s_type
                slides_with_content.append(content)
        await db.presentations.update_one(
            {"id": pres_id},
            {"$set": {"status": "generating_layout", "step": 3, "slides": slides_with_content}}
        )

        # Step 3: Layout
        logger.info(f"[{pres_id}] Step 3: Generating layouts...")
        slides_with_layout = []
        for s in slides_with_content:
            laid_out = await llm_layout_for_slide(s)
            slides_with_layout.append(laid_out)
        await db.presentations.update_one(
            {"id": pres_id},
            {"$set": {"status": "rendering", "step": 4, "slides": slides_with_layout}}
        )

        # Step 4: Render PPTX & PDF
        logger.info(f"[{pres_id}] Step 4: Rendering files...")
        pptx_path = render_pptx(pres_id, slides_with_layout, pres_title)
        pdf_path = render_pdf(pres_id, slides_with_layout, pres_title)

        await db.presentations.update_one(
            {"id": pres_id},
            {"$set": {
                "status": "completed",
                "step": 5,
                "slides": slides_with_layout,
                "pptx_path": pptx_path,
                "pdf_path": pdf_path,
            }}
        )
        logger.info(f"[{pres_id}] Generation complete!")
    except Exception as e:
        logger.error(f"[{pres_id}] Generation failed: {e}")
        await db.presentations.update_one(
            {"id": pres_id},
            {"$set": {"status": "failed", "error": str(e)}}
        )

# ── App setup ──
app = FastAPI()
api_router = APIRouter(prefix="/api")

# ── Pydantic models ──
class RegisterRequest(BaseModel):
    email: str
    password: str
    name: str

class LoginRequest(BaseModel):
    email: str
    password: str

class GenerateRequest(BaseModel):
    prompt: str

# ── Auth routes ──
@api_router.post("/auth/register")
async def register(body: RegisterRequest, response: Response):
    email = body.email.lower().strip()
    if await db.users.find_one({"email": email}):
        raise HTTPException(400, "Email already registered")
    if len(body.password) < 6:
        raise HTTPException(400, "Password must be at least 6 characters")
    hashed = hash_password(body.password)
    user_doc = {
        "email": email,
        "password_hash": hashed,
        "name": body.name.strip(),
        "role": "user",
        "created_at": datetime.now(timezone.utc).isoformat(),
    }
    result = await db.users.insert_one(user_doc)
    user_id = str(result.inserted_id)
    access = create_access_token(user_id, email)
    refresh = create_refresh_token(user_id)
    set_auth_cookies(response, access, refresh)
    return {"_id": user_id, "email": email, "name": body.name.strip(), "role": "user"}

@api_router.post("/auth/login")
async def login(body: LoginRequest, request: Request, response: Response):
    email = body.email.lower().strip()
    ip = request.client.host if request.client else "unknown"
    await check_brute_force(ip, email)
    user = await db.users.find_one({"email": email})
    if not user or not verify_password(body.password, user["password_hash"]):
        await record_failed_attempt(ip, email)
        raise HTTPException(401, "Invalid email or password")
    await clear_failed_attempts(ip, email)
    user_id = str(user["_id"])
    access = create_access_token(user_id, email)
    refresh = create_refresh_token(user_id)
    set_auth_cookies(response, access, refresh)
    return {"_id": user_id, "email": email, "name": user.get("name", ""), "role": user.get("role", "user")}

@api_router.post("/auth/logout")
async def logout(response: Response):
    response.delete_cookie("access_token", path="/")
    response.delete_cookie("refresh_token", path="/")
    return {"message": "Logged out"}

@api_router.get("/auth/me")
async def get_me(request: Request):
    user = await get_current_user(request)
    return user

@api_router.post("/auth/refresh")
async def refresh_token(request: Request, response: Response):
    token = request.cookies.get("refresh_token")
    if not token:
        raise HTTPException(401, "No refresh token")
    try:
        payload = pyjwt.decode(token, get_jwt_secret(), algorithms=[JWT_ALGORITHM])
        if payload.get("type") != "refresh":
            raise HTTPException(401, "Invalid token type")
        user = await db.users.find_one({"_id": ObjectId(payload["sub"])})
        if not user:
            raise HTTPException(401, "User not found")
        user_id = str(user["_id"])
        access = create_access_token(user_id, user["email"])
        set_auth_cookies(response, access, token)
        return {"message": "Token refreshed"}
    except pyjwt.ExpiredSignatureError:
        raise HTTPException(401, "Refresh token expired")
    except pyjwt.InvalidTokenError:
        raise HTTPException(401, "Invalid refresh token")

# ── Presentation routes ──
@api_router.post("/presentations/generate")
async def start_generation(body: GenerateRequest, request: Request):
    user = await get_current_user(request)
    prompt = body.prompt.strip()
    if not prompt:
        raise HTTPException(400, "Prompt cannot be empty")
    pres_id = str(uuid.uuid4())
    doc = {
        "id": pres_id,
        "user_id": user["_id"],
        "prompt": prompt,
        "title": "",
        "slides": [],
        "status": "orchestrating",
        "step": 1,
        "total_steps": 5,
        "created_at": datetime.now(timezone.utc).isoformat(),
        "pptx_path": "",
        "pdf_path": "",
        "error": "",
    }
    await db.presentations.insert_one(doc)
    asyncio.create_task(generate_presentation_task(pres_id, prompt))
    return {"id": pres_id, "status": "orchestrating"}

@api_router.get("/presentations")
async def list_presentations(request: Request):
    user = await get_current_user(request)
    items = await db.presentations.find(
        {"user_id": user["_id"]},
        {"_id": 0}
    ).sort("created_at", -1).to_list(100)
    return items

@api_router.get("/presentations/{pres_id}")
async def get_presentation(pres_id: str, request: Request):
    user = await get_current_user(request)
    doc = await db.presentations.find_one(
        {"id": pres_id, "user_id": user["_id"]},
        {"_id": 0}
    )
    if not doc:
        raise HTTPException(404, "Presentation not found")
    return doc

@api_router.delete("/presentations/{pres_id}")
async def delete_presentation(pres_id: str, request: Request):
    user = await get_current_user(request)
    doc = await db.presentations.find_one({"id": pres_id, "user_id": user["_id"]})
    if not doc:
        raise HTTPException(404, "Presentation not found")
    # Delete files
    for key in ["pptx_path", "pdf_path"]:
        fpath = doc.get(key, "")
        if fpath and Path(fpath).exists():
            Path(fpath).unlink()
    await db.presentations.delete_one({"id": pres_id})
    return {"message": "Deleted"}

@api_router.get("/presentations/{pres_id}/download/pptx")
async def download_pptx(pres_id: str, request: Request):
    user = await get_current_user(request)
    doc = await db.presentations.find_one({"id": pres_id, "user_id": user["_id"]}, {"_id": 0})
    if not doc:
        raise HTTPException(404, "Presentation not found")
    fpath = doc.get("pptx_path", "")
    if not fpath or not Path(fpath).exists():
        raise HTTPException(404, "PPTX file not found")
    safe_title = "".join(c for c in doc.get("title", "presentation") if c.isalnum() or c in " -_")[:50]
    return FileResponse(fpath, media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation", filename=f"{safe_title}.pptx")

@api_router.get("/presentations/{pres_id}/download/pdf")
async def download_pdf(pres_id: str, request: Request):
    user = await get_current_user(request)
    doc = await db.presentations.find_one({"id": pres_id, "user_id": user["_id"]}, {"_id": 0})
    if not doc:
        raise HTTPException(404, "Presentation not found")
    fpath = doc.get("pdf_path", "")
    if not fpath or not Path(fpath).exists():
        raise HTTPException(404, "PDF file not found")
    safe_title = "".join(c for c in doc.get("title", "presentation") if c.isalnum() or c in " -_")[:50]
    return FileResponse(fpath, media_type="application/pdf", filename=f"{safe_title}.pdf")

# ── Include router ──
app.include_router(api_router)

# ── CORS ──
cors_origins = [
    os.environ.get("FRONTEND_URL", "http://localhost:3000"),
    "http://localhost:3000",
]
# Add any CORS_ORIGINS from env
extra = os.environ.get("CORS_ORIGINS", "")
if extra and extra != "*":
    cors_origins.extend(extra.split(","))

app.add_middleware(
    CORSMiddleware,
    allow_origins=cors_origins,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# ── Startup ──
@app.on_event("startup")
async def startup():
    await db.users.create_index("email", unique=True)
    await db.login_attempts.create_index("identifier")
    # Seed admin
    admin_email = os.environ.get("ADMIN_EMAIL", "admin@example.com")
    admin_password = os.environ.get("ADMIN_PASSWORD", "admin123")
    existing = await db.users.find_one({"email": admin_email})
    if not existing:
        await db.users.insert_one({
            "email": admin_email,
            "password_hash": hash_password(admin_password),
            "name": "Admin",
            "role": "admin",
            "created_at": datetime.now(timezone.utc).isoformat(),
        })
        logger.info(f"Admin seeded: {admin_email}")
    elif not verify_password(admin_password, existing["password_hash"]):
        await db.users.update_one({"email": admin_email}, {"$set": {"password_hash": hash_password(admin_password)}})
        logger.info("Admin password updated")
    # Write test credentials
    Path("/app/memory").mkdir(exist_ok=True)
    with open("/app/memory/test_credentials.md", "w") as f:
        f.write(f"# Test Credentials\n\n## Admin\n- Email: {admin_email}\n- Password: {admin_password}\n- Role: admin\n\n## Auth Endpoints\n- POST /api/auth/register\n- POST /api/auth/login\n- POST /api/auth/logout\n- GET /api/auth/me\n- POST /api/auth/refresh\n\n## Presentation Endpoints\n- POST /api/presentations/generate\n- GET /api/presentations\n- GET /api/presentations/:id\n- DELETE /api/presentations/:id\n- GET /api/presentations/:id/download/pptx\n- GET /api/presentations/:id/download/pdf\n")

@app.on_event("shutdown")
async def shutdown():
    client.close()
